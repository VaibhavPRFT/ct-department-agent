#!/usr/bin/env python3
"""
import_webinars.py — turn the Webinar Series Tracker workbook into data/webinars.json.

Usage (from the repo root):
    pip install openpyxl
    python agent/import_webinars.py "<path to tracker .xlsx>" [--weeks "Wk01 Sep 28,Wk02 Oct 05"]

Reads the Command Center, Sessions, Speakers and weekly (WkNN ...) tabs.
By default every weekly tab is imported; pass --weeks to limit which ones are
published. Review notes already in data/webinars.json (weeks[].review) are
kept, so re-running the import after a status update won't wipe them.
"""
import argparse
import datetime as dt
import json
import os
import re
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "data", "webinars.json")
TEAMS = ["Sales", "Marketing", "Practice", "Partnership"]


def iso(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    return v


def rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def header_index(rs, first):
    for i, r in enumerate(rs):
        if r and r[0] == first:
            return i
    raise ValueError(f"header '{first}' not found")


def read_sessions(wb):
    rs = rows(wb["Sessions"])
    h = header_index(rs, "#")
    cols = [str(c).strip() if c else "" for c in rs[h]]
    out = []
    for r in rs[h + 1:]:
        if not isinstance(r[0], (int, float)):
            continue
        d = dict(zip(cols, r))
        out.append({
            "n": int(d["#"]),
            "date": iso(d["Date"]),
            "practice": d["Practice"],
            "title": d["Session Title"],
            "format": d["Format"],
            "wave": d["Campaign Wave"],
            "offer": d["Complimentary Offer (CTA)"],
            "owner": d["Sales Owner"],
            "regTarget": d.get("Reg Target") or 0,
            "registered": d.get("Registered") or 0,
            "attended": d.get("Attended") or 0,
            "mqls": d.get("MQLs") or 0,
            "assessments": d.get("Assessments Booked") or 0,
            "health": d.get("Health") or "On Track",
        })
    return out


def read_speakers(wb):
    rs = rows(wb["Speakers"])
    h = header_index(rs, "#")
    out = {}
    for r in rs[h + 1:]:
        if not isinstance(r[0], (int, float)):
            continue
        out[int(r[0])] = {
            "executive": r[5], "practice": r[7], "guest": r[9],
            "bios": r[11], "trapsValidated": r[12], "deck": r[13],
            "dryRun": iso(r[14]), "dryRunDone": r[15], "readiness": r[16] or 0,
        }
    return out


def read_week(ws):
    rs = rows(ws)
    title = rs[0][0].strip()
    meta = {}
    for r in rs[1:6]:
        if r[0] in ("ANCHOR", "ALSO IN MOTION", "CONVERSION PATH"):
            meta[r[0]] = next((c for c in r[1:] if c), "")
    h = header_index(rs, "Task ID")
    tasks = []
    for r in rs[h + 1:]:
        if not r[0]:
            continue
        tasks.append({
            "id": r[0], "day": r[1], "date": iso(r[2]), "team": r[3],
            "owner": r[4], "task": r[5], "session": r[6],
            "deliverable": r[7], "status": r[8] or "Not Started", "notes": r[9],
        })
    m = re.search(r"Mon (\w+ \d+) to Fri (\w+ \d+), (\d{4})", title)
    progress = {
        t: {
            "total": sum(1 for x in tasks if x["team"] == t),
            "done": sum(1 for x in tasks if x["team"] == t and x["status"] == "Done"),
        }
        for t in TEAMS
    }
    return {
        "tab": ws.title,
        "title": title,
        "range": f"{m.group(1)} – {m.group(2)}, {m.group(3)}" if m else title,
        "anchor": meta.get("ANCHOR", ""),
        "alsoInMotion": meta.get("ALSO IN MOTION", "").strip(),
        "progress": progress,
        "tasks": tasks,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--weeks", help="comma-separated weekly tab names to publish")
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.xlsx, data_only=True)
    week_tabs = [n for n in wb.sheetnames if re.match(r"Wk\d+ ", n)]
    if a.weeks:
        wanted = [w.strip() for w in a.weeks.split(",")]
        missing = [w for w in wanted if w not in week_tabs]
        if missing:
            sys.exit(f"weekly tabs not found: {missing}")
        week_tabs = wanted

    prev = {}
    if os.path.exists(OUT):
        with open(OUT, encoding="utf-8") as f:
            old = json.load(f)
        prev = {w["tab"]: w.get("review", []) for w in old.get("weeks", [])}
        meta = old.get("meta", {})
    else:
        meta = {}

    cc = rows(wb["Command Center"])
    summary = next((r[0] for r in cc if r[0] and str(r[0]).startswith("12 ")), "")
    sessions = read_sessions(wb)
    speakers = read_speakers(wb)
    for s in sessions:
        s["speakers"] = speakers.get(s["n"], {})

    weeks = []
    for n in week_tabs:
        w = read_week(wb[n])
        w["review"] = prev.get(n, [])
        weeks.append(w)

    meta.setdefault("eyebrow", "Royal Cyber · commercetools + MetafyAI")
    meta.setdefault("title", "Webinar Series Tracker")
    meta["summary"] = summary
    meta["updated"] = "Updated " + dt.date.today().strftime("%b %d, %Y").replace(" 0", " ")

    data = {"meta": meta, "sessions": sessions, "weeks": weeks}
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
        f.write("\n")
    print(f"wrote {OUT}: {len(sessions)} sessions, weeks {[w['tab'] for w in weeks]}")


if __name__ == "__main__":
    main()
